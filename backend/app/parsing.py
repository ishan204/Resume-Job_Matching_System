"""File-to-text extraction and structured parsing, built only on the hybrid's own parsers
(ml/features/*). Uploaded bytes are read in memory and never written to disk or logged."""
import io
import re

from ml.features.education import LEVELS, degree_levels
from ml.features.responsibilities import candidate_evidence
from ml.features.sections import labelled_segments
from ml.models.hybrid import HybridMatcher
from ml.models.text import normalize

ALLOWED = {".pdf": "application/pdf", ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
           ".txt": "text/plain"}


class ParseError(ValueError):
    """Raised for unsupported or unreadable uploads; message is safe to show to users."""


def extract_text(filename: str, data: bytes) -> str:
    ext = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if ext not in ALLOWED:
        raise ParseError(f"Unsupported file type '{ext or filename}'. Allowed: PDF, DOCX, TXT.")
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
        elif ext == ".docx":
            from docx import Document
            doc = Document(io.BytesIO(data))
            text = "\n".join([p.text for p in doc.paragraphs] +
                             [c.text for t in doc.tables for row in t.rows for c in row.cells])
        else:
            try:
                text = data.decode("utf-8")
            except UnicodeDecodeError:
                text = data.decode("latin-1")
    except Exception as exc:  # malformed file: report safely, never the parser's internals
        raise ParseError(f"Could not read the {ext[1:].upper()} file; it may be corrupted or encrypted.") from exc
    if not text.strip():
        raise ParseError("No text could be extracted from the file (scanned PDFs without a text layer "
                         "are not supported).")
    return text


def _section_sentences(text: str, kind: str) -> list[str]:
    return [text[s:e] for s, e, section in labelled_segments(text, "resume")
            if section == kind and len(text[s:e].split()) >= 3]


def parse_resume(parser: HybridMatcher, raw: str, include_text: bool = False) -> dict:
    t = normalize(raw)
    r = parser.parse_resume(raw)
    cats = parser.extractor.category
    education = []
    for sentence in _section_sentences(t, "education"):
        for level in degree_levels(sentence):
            education.append({"level": LEVELS[level], "evidence": sentence})
    out = {
        "skills": [{"skill": s, "category": cats[s], "evidence": ev} for s, ev in sorted(r["skills"].items())],
        "education": education,
        "highest_degree": LEVELS.get(r["education"]["highest_level"]),
        "experience": {
            "total_years": r["experience"]["total_years"],
            "source": r["experience"]["source"],
            "date_ranges_found": r["experience"]["date_ranges"],
            "stated_years": r["experience"]["stated_years"],
            "statements": candidate_evidence(t),
        },
        "projects": _section_sentences(t, "projects"),
        "certifications": sorted({s for s in r["skills"] if cats[s] == "certification"} |
                                 set(_section_sentences(t, "certifications"))),
        "notes": ["Experience years are estimated from employment date ranges (overlaps merged; "
                  "'Present' = latest explicit date) and may undercount current roles.",
                  "Skills are matched against a fixed 178-skill taxonomy; soft skills are not extracted."],
    }
    if include_text:
        out["text"] = raw
    return out


def _title(raw: str) -> str | None:
    """Heuristic only (not used by any model): the first non-empty line, if it is short and not a sentence."""
    first = next((line.strip(" :-	") for line in raw.splitlines() if line.strip()), "")
    return first if 1 <= len(first.split()) <= 8 and not re.search(r"[.!?]$", first) else None


def parse_job(parser: HybridMatcher, raw: str) -> dict:
    t = normalize(raw)
    j = parser.parse_job(raw)
    by = lambda status: [{"skill": s, "evidence": v["evidence"]} for s, v in sorted(j["skills"].items())
                         if v["status"] == status]
    e, ed = j["experience"], j["education"]
    return {
        "title": _title(raw),
        "required_skills": by("required"),
        "preferred_skills": by("preferred"),
        "uncertain_skills": by("uncertain"),
        "experience_requirement": {"min_years": e["required_min_years"], "max_years": e["required_max_years"],
                                   "preferred_years": e["preferred_years"], "evidence": e["evidence"]},
        "education": {"required_level": LEVELS.get(ed["required_level"]),
                      "preferred_level": LEVELS.get(ed["preferred_level"])},
        "responsibilities": j["responsibilities"],
        "responsibilities_from_section": j["responsibilities_from_section"],
        "notes": ["'uncertain' skills are mentioned without required/preferred wording or section.",
                  "The title is a heuristic guess from the first line and is not used for matching."],
    }
